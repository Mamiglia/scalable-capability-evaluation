"""Results table in the shared sheet's layout, filled from analysis/results/forecasts.csv.

    PYTHONPATH=src uv run python scripts/analyses/results_table.py [out.xlsx]

Rows: baselines and reference lines, the transcript conditions, then each judge condition as
one row per judge (core condition) followed by BEST (best judge on references), mean over
judges and pooled judges. Columns: Protocol, Description, Model, one per domain x semester
(two header rows, domain merged across its semesters), and a pooled column. Formatting copies
the sheet: purple header for the label columns, blue for the task columns, lavender / light
blue body fills, per-judge values bold, label cells merged over each judge block.
"""

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analysis import stats
from reference import tables

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("analysis/results/tables/results_table.xlsx")
DOMS = [("science", ["2024-I", "2024-II", "2025-I", "2025-II", "2026-I"]), ("math", ["2024-I", "2024-II", "2025-I", "2025-II", "2026-I"]),
        ("coding", ["2025-II", "2026-I"]), ("cyber", ["2025-II", "2026-I"])]
COLS = [(d, s) for d, ss in DOMS for s in ss]
HEAD_LABEL, HEAD_TASK, BODY_LABEL, BODY_VALUE = "FF674EA7", "FF3C78D8", "FFD9D2E9", "FFC9DAF8"


def mse(g: pd.DataFrame) -> float:
    return float(((g.theta_hat - g.theta) ** 2).mean()) if len(g) else np.nan


def build_rows(f: pd.DataFrame) -> list[dict]:
    """[{protocol, description, model, values (per COLS + pooled), bold}] in sheet order."""
    core, summ = stats.core_condition(f), stats.judge_summaries(f)
    rows = []

    def add(protocol, description, model, sub, bold=False, block=None):
        rows.append(dict(protocol=protocol, description=description, model=model, bold=bold, block=block,
                         values=[mse(sub[(sub.domain == d) & (sub.semester == s)]) for d, s in COLS] + [mse(sub)]))
    frozen = {}
    for d, ss in DOMS:
        m = tables.load(d)[0]
        top = tables.fit_settings(d)["top"]
        for s in ss:
            refs = m.loc[m[f"{s}_role"] == "reference", f"{s}_theta"]
            tg = m.loc[m[f"{s}_role"] == "target", f"{s}_theta"]
            frozen[(d, s)] = (((tg - refs[top[s]].mean()) ** 2).mean(), ((tg - refs.max()) ** 2).mean(), len(tg))

    def const(i, protocol, description):
        vals = [frozen[c][i] for c in COLS]
        rows.append(dict(protocol=protocol, description=description, model="-", bold=False, block=None,
                         values=vals + [float(np.average(vals, weights=[frozen[c][2] for c in COLS]))]))
    add("Prediction from saturated items", "the references' mean theta* (accuracy on the saturated items is 1 for every model)", "-", f[f.protocol == "easy_items"])
    const(0, "Top-3 reference mean", "mean theta* of the three strongest references")
    const(1, "Best reference", "theta* of the strongest reference")
    add("Accuracy oracle (HARD items)", "OLS of theta* on all-repeat accuracy over the task's HARD items, fit on the references", "-", f[f.protocol == "accuracy_hard"])
    add("Length", "ridge on 3 length features, penalty by leave-one-out on the references", "-", f[f.protocol == "glm_length"])
    add("Length + prose features", "ridge on 3 length + 9 regex features, penalty by leave-one-out", "-", f[f.protocol == "glm_regex"])
    for v in ("full", "answer"):
        add(f"Embedding ({v})", "ridge on the mean Qwen3-Embedding-8B vector, penalty by leave-one-out", "-", f[(f.protocol == "embedding") & (f.variant == v)])
    blocks = [("direct", "raw", "Direct scoring (raw ratings)", "1-10 ratings, OLS of theta* on the mean rating; reasoning shown where judged"),
              ("direct", "item_z", "Direct scoring (item z-scored)", "ratings z-scored per item on the references, then OLS"),
              ("pairwise", "linear", "Pairwise judge (linear calibration)", "per-bench slope/noise on reference pairs, k>0 only, sampling-noise floor, references' N(mean, var) prior; reasoning shown where judged"),
              ("pairwise", "bt", "Pairwise judge (pinned BT)", "plain Bradley-Terry with the references pinned at theta*"),
              ("pairwise", "calibrated", "Pairwise judge (calibrated BT)", "calibrated BT on reference pairs + leave-one-out-error shrinkage to the frontier prior")]
    for protocol, variant, title, desc in blocks:
        sub = core[(core.protocol == protocol) & (core.variant == variant)]
        for judge in sorted(sub.judge.unique()):
            add(title, desc, judge, sub[sub.judge == judge], bold=True, block=title)
        for label in (stats.BEST, stats.MEAN):
            add(f"{title} (summary)", "", label, summ[(summ.protocol == protocol) & (summ.variant == variant) & (summ.judge == label)], block=f"{title} (summary)")
        ens = f[(f.protocol == f"{protocol}_ensemble") & (f.variant == variant)]
        if len(ens):
            add(f"{title} (summary)", "", "pooled judges (ensemble)", ens, block=f"{title} (summary)")
    return rows


def write(rows: list[dict], path: Path) -> None:
    wb = Workbook(); ws = wb.active; ws.title = "results"
    thin = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A1"], ws["B1"], ws["C1"] = "Protocol", "Description", "Model"
    ws.merge_cells("A1:A2"); ws.merge_cells("B1:B2"); ws.merge_cells("C1:C2")
    col = 4
    for d, ss in DOMS:
        ws.cell(1, col, d.upper()); ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(ss) - 1)
        for i, s in enumerate(ss):
            ws.cell(2, col + i, s)
        col += len(ss)
    ws.cell(1, col, "POOLED"); ws.cell(2, col, "all targets")
    last = col
    for c in range(1, last + 1):
        for r in (1, 2):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEAD_LABEL if c <= 3 else HEAD_TASK)
            cell.alignment = thin
    r = 3
    starts: dict[str, int] = {}
    for row in rows:
        ws.cell(r, 1, row["protocol"]); ws.cell(r, 2, row["description"]); ws.cell(r, 3, row["model"])
        for c in (1, 2, 3):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=BODY_LABEL)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 3).font = Font(bold=row["model"] in (stats.BEST, stats.MEAN, "pooled judges (ensemble)"))
        for i, v in enumerate(row["values"]):
            cell = ws.cell(r, 4 + i, None if not np.isfinite(v) else round(float(v), 2))
            cell.fill = PatternFill("solid", fgColor=BODY_VALUE)
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
            cell.font = Font(bold=row["bold"])
        if row["block"]:
            starts.setdefault(row["block"], r)
            if r > starts[row["block"]]:
                pass
        r += 1
    # merge the label cells over each judge block (protocol and description)
    for block, start in starts.items():
        end = max(i for i, row in enumerate(rows, 3) if row["block"] == block)
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
            ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)
            ws.cell(start, 1).alignment = Alignment(vertical="top", wrap_text=True)
    for c, w in zip("ABC", (30, 34, 27)):
        ws.column_dimensions[c].width = w
    for c in range(4, last + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = "D3"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    forecasts = pd.read_csv("analysis/results/forecasts.csv")
    table = build_rows(forecasts)
    write(table, OUT)
    print(f"wrote {OUT}: {len(table)} rows")
